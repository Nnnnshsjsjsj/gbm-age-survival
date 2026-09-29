# -*- coding: utf-8 -*-
import re, sys, copy, importlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

lang = sys.argv[1]
C = importlib.import_module(f"content_{lang}")
out = sys.argv[2]

doc = Document("template.docx")
body = doc.element.body
for el in list(body):
    if el.tag != qn("w:sectPr"):
        body.remove(el)

styles = doc.styles
if "Normal" not in [s.name for s in styles]:
    st = styles.add_style("Normal", WD_STYLE_TYPE.PARAGRAPH)
    st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.5
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # make it the default paragraph style
    st.element.set(qn("w:default"), "1")
def S(name):
    for s_ in styles:
        if s_.name == name: return s_
    raise KeyError(name)
normal = S("Normal")
for hn in ["Heading 1", "Heading 2", "Heading 3"]:
    h = S(hn)
    h.paragraph_format.space_before = Pt(18 if hn == "Heading 1" else 12)
    h.paragraph_format.space_after = Pt(6)
    h.paragraph_format.keep_with_next = True
    h.font.bold = True
    h.font.name = "Times New Roman"
    rpr = h.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), "Times New Roman")
    for k in ("w:lang",):
        pass
# fix eastAsia/cs fonts on Normal too (Cyrillic renders from hAnsi, fine)

# ---------- citation resolution ----------
order = []
def cite(m):
    keys = [k.strip().lstrip("@") for k in m.group(1).split(";")]
    nums = []
    for k in keys:
        if k not in C.REFS:
            raise SystemExit(f"unknown ref key {k}")
        if k not in order:
            order.append(k)
        nums.append(order.index(k) + 1)
    nums = sorted(set(nums))
    # compress runs
    parts, i = [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else (f"{nums[i]}, {nums[j]}" if j > i else f"{nums[i]}"))
        i = j + 1
    return "[" + ", ".join(parts) + "]"

TOKEN = re.compile(r"\[((?:@\w+;?\s*)+)\]")
def resolve(text):
    return TOKEN.sub(cite, text)

# ---------- helpers ----------
def para(text="", style=None, align=None, size=None, bold=None, italic=None, space_after=None, first_indent=None, keep=False):
    p = doc.add_paragraph(style=(S(style) if style else None))
    if text:
        r = p.add_run(text)
        if size: r.font.size = Pt(size)
        if bold is not None: r.bold = bold
        if italic is not None: r.italic = italic
    if align is not None: p.alignment = align
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if first_indent is not None: p.paragraph_format.first_line_indent = Cm(first_indent)
    if keep: p.paragraph_format.keep_with_next = True
    return p

def page_break():
    p = doc.add_paragraph(); p.add_run().add_break(WD_BREAK.PAGE)

def set_cell_borders(tbl):
    tblPr = tbl._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4"); e.set(qn("w:space"), "0"); e.set(qn("w:color"), "808080")
        borders.append(e)
    tblPr.append(borders)

def table(caption, header, rows):
    para(caption, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4, keep=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_borders(t)
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(10)
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "E7EEF6"); c._tc.get_or_add_tcPr().append(shd)
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = t.rows[i].cells[j]; c.text = ""
            r = c.paragraphs[0].add_run(resolve(val)); r.font.size = Pt(10)
    # Wide tables (7+ columns): one font step smaller so headers do not break mid-word.
    if len(header) >= 7:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    for r in p.runs: r.font.size = Pt(8.5)
    nrows = len(t.rows)
    for ri, row in enumerate(t.rows):
        trPr_ = row._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr_.append(cs)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.0
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.first_line_indent = Cm(0)
                if ri < nrows - 1 and nrows <= 16:
                    p.paragraph_format.keep_with_next = True
    # column widths proportional to content length
    import math
    ncol = len(header)
    lens = [max([len(header[j])] + [len(r[j]) for r in rows]) for j in range(ncol)]
    w = [math.sqrt(max(l, 4)) for l in lens]
    total = Cm(16.0); s_ = sum(w)
    t.autofit = False
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:w"), str(int(16.0 / 2.54 * 1440))); tblW.set(qn("w:type"), "dxa"); t._tbl.tblPr.append(tblW)
    widths = [int(total * w[j] / s_) for j in range(ncol)]
    # no column narrower than the longest single word in it (plus padding), so words never split across lines
    import re as _re
    minw = [Cm(0.6 + 0.215 * max(len(wd) for cell in [header[j]] + [r[j] for r in rows] for wd in _re.split(r"[\s/]+", str(cell)) or [""])) for j in range(ncol)]
    minw = [min(mw, Cm(4.5)) for mw in minw]
    free = total - sum(minw)
    if free >= 0:
        widths = [mw + free * w[j] / s_ for j, mw in enumerate(minw)]
    else:
        k = total / sum(minw); widths = [mw * k for mw in minw]
    widths = [int(wd) for wd in widths]
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); t._tbl.tblPr.append(lay)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int(widths[j] / 635)))  # EMU -> twips
    for j in range(ncol):
        for row in t.rows:
            row.cells[j].width = widths[j]
    # repeat header row
    trPr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true"); trPr.append(th)
    para("", space_after=6)

def figure(key, caption):
    import glob
    path = sorted(glob.glob(f"figs/{key}_{lang}.*"))[0]
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    width = Inches(6.2) if "flow" in key else Inches(6.0)
    p.add_run().add_picture(path, width=width)
    cp = para(resolve(caption), size=11, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    cp.paragraph_format.line_spacing = 1.15

def numbered(items):
    for i, it in enumerate(items, 1):
        # items that carry their own label (H1…, Г1…) are not numbered again
        own = it.split(" ")[0].rstrip(".") if it[:2] in ("H1","H2","H3","H4","Г1","Г2","Г3","Г4") else None
        p = para((it if own else f"{i}. " + it) if False else (resolve(it) if own else f"{i}. " + resolve(it)), style="List Paragraph")
        p.paragraph_format.left_indent = Cm(0.75); p.paragraph_format.first_line_indent = Cm(-0.75)
        p.paragraph_format.space_after = Pt(4)

def render(items):
    for it in items:
        kind = it[0]
        if kind == "h1":
            if it[1] not in (C.LABELS.get("abstract"),):
                page_break()
            para(it[1], style="Heading 1")
        elif kind == "h2": para(it[1], style="Heading 2")
        elif kind == "h3": para(it[1], style="Heading 3")
        elif kind == "p": para(resolve(it[1]), first_indent=1.0)
        elif kind == "table": table(*it[1:])
        elif kind == "fig": figure(it[1], it[2])
        elif kind == "num": numbered(it[1])
        else: raise SystemExit(kind)

# ---------- cover ----------
for i, line in enumerate(C.COVER):
    if line == C.TITLE:
        p = para(line, style="Title", align=WD_ALIGN_PARAGRAPH.CENTER); p.paragraph_format.space_before = Pt(60); p.paragraph_format.space_after = Pt(60)
        for r in p.runs: r.font.size = Pt(22); r.bold = True
    elif line == "":
        para("")
    else:
        p = para(line, align=WD_ALIGN_PARAGRAPH.CENTER, size=13)
        if line.startswith("Author") or line.startswith("Автор") or line.startswith("Supervisor") or line.startswith("Научный") or line.startswith("Advisor") or line.startswith("Консультант"):
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if i == 2:
            p.paragraph_format.space_before = Pt(120)
            for r in p.runs: r.bold = True
        if i == len(C.COVER) - 1:
            p.paragraph_format.space_before = Pt(160)
page_break()

# ---------- TOC ----------
para(C.LABELS["toc"], style="Heading 1")
p = doc.add_paragraph()
r = p.add_run()
fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin"); r._r.append(fld)
r2 = p.add_run(); instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = 'TOC \\o "1-3" \\h \\z \\u'; r2._r.append(instr)
r3 = p.add_run(); fld2 = OxmlElement("w:fldChar"); fld2.set(qn("w:fldCharType"), "separate"); r3._r.append(fld2)
r4 = p.add_run(C.LABELS["toc_note"]); r4.italic = True; r4.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
r5 = p.add_run(); fld3 = OxmlElement("w:fldChar"); fld3.set(qn("w:fldCharType"), "end"); r5._r.append(fld3)
page_break()

# ---------- abstract ----------
para(C.LABELS["abstract"], style="Heading 1")
para(C.ABSTRACT, first_indent=1.0)
para(C.LABELS["keywords"], style="Heading 2")
para(C.KEYWORDS)

# ---------- body ----------
render(C.BODY)

# ---------- references ----------
page_break()
para(C.LABELS["references"], style="Heading 1")
for i, k in enumerate(order, 1):
    p = para(f"{i}. " + C.REFS[k])
    p.paragraph_format.left_indent = Cm(0.9); p.paragraph_format.first_line_indent = Cm(-0.9)
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.15
    for r in p.runs: r.font.size = Pt(11)
unused = [k for k in C.REFS if k not in order]
if unused: print("UNUSED REFS:", unused)

# ---------- appendices ----------
render(C.APPENDICES)

# document-level: update fields prompt
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true"); settings.append(uf)

# page numbers in footer
sec = doc.sections[0]
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run(); f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin"); r._r.append(f1)
r = fp.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = "PAGE"; r._r.append(it)
r = fp.add_run(); f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end"); r._r.append(f2)

doc.save(out)
print("saved", out, "refs:", len(order))
